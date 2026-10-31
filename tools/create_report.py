"""Generador histórico del informe inicial.

El informe vigente está en docs/Informe arquitectura Ajedrez Minecraft.docx.
Este script conserva la versión anterior como referencia y genera una salida
separada en build/. No reproduce las revisiones UML ni las etapas 5 y 6.
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

doc=Document()
section=doc.sections[0]
section.page_height=Inches(11.7);section.page_width=Inches(8.3)
section.top_margin=section.bottom_margin=Inches(.73)
section.left_margin=section.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.07
for name,size in [('Title',27),('Heading 1',19),('Heading 2',13)]:
 st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(9)
doc.styles['Subtitle'].font.size=Pt(12)
doc.core_properties.title='Arquitectura del ajedrez de Minecraft'
doc.core_properties.subject='Clases del núcleo, puertos de entrada y salida y adaptadores'
doc.core_properties.author='Proyecto Ajedrez TPO'
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Ajedrez TPO  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for run in footer.runs:run.font.size=Pt(8)

def p(text):return doc.add_paragraph(text)
def h(text,level=1):return doc.add_heading(text,level=level)
def page(title):doc.add_page_break();h(title)
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for cell,width in zip(t.columns,widths):cell.width=Inches(width)
 pr=t._tbl.tblPr
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
 pr.append(borders)
 for i,text in enumerate(headers):t.rows[0].cells[i].text=text
 for row in rows:
  cells=t.add_row().cells
  for cell,text in zip(cells,row):cell.text=text
 for ri,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();keep=OxmlElement('w:cantSplit');trpr.append(keep)
  if ri==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcp=cell._tc.get_or_add_tcPr()
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'374151' if ri==0 else ('F3F4F6' if ri%2==0 else 'FFFFFF'));tcp.append(shade)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    node=OxmlElement('w:'+edge);node.set(qn('w:w'),'90');node.set(qn('w:type'),'dxa');margins.append(node)
   tcp.append(margins)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1.02
    for run in para.runs:
     run.font.size=Pt(9.5)
     if ri==0:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t

doc.add_paragraph('Arquitectura del ajedrez de Minecraft','Title')
doc.add_paragraph('Clases del núcleo, puertos de entrada y salida y adaptadores','Subtitle')
p('Ingeniería de Software · Trabajo práctico de ajedrez · 12 de octubre de 2026')
h('Propósito del informe')
p('El proyecto separa las reglas del ajedrez de la interfaz gráfica y del almacenamiento. Esta organización permite verificar movimientos y jaques sin abrir un navegador, cambiar la presentación sin reescribir las reglas y conservar una partida mediante un contrato de guardado.')
p('El informe describe la implementación del proyecto: las clases que forman el núcleo, las interfaces que permiten entrar y salir de él, y los adaptadores que conectan el juego con el usuario y el disco. La distinción se explica con una jugada concreta y con las funciones incorporadas de puntaje, cronómetro e inteligencia artificial.')
h('Organización general')
table(['Parte','Contenido y responsabilidad'],[
('Núcleo de reglas','Paquete core: tablero, piezas, movimientos, jaque, turnos, historial y resultado. No conoce HTML ni archivos.'),
('Casos de uso de la sesión','MatchSession coordina modo local o contra IA, cronómetro, guardado y recuperación. Se apoya en contratos.'),
('Adaptadores de entrada','WebServer traduce solicitudes HTTP. ConsoleUI traduce comandos de terminal para el juego básico.'),
('Adaptador de salida','FileGameRepository transforma los datos de SavedMatch en un archivo local y los recupera.'),
('Presentación','HTML, CSS y JavaScript muestran personajes, destinos, puntajes y tutorial. Java confirma las jugadas.')],[1.75,4.95])
h('Dirección de las dependencias',2)
p('La web y la consola dependen de interfaces del núcleo. El guardado implementa una interfaz definida por el núcleo. Así, las reglas no necesitan importar clases de HTTP, del navegador ni del sistema de archivos. StandardGame y el arranque de WebServer crean los objetos y conectan sus dependencias por constructor.')
p('Es un monolito pequeño con una separación de puertos y adaptadores. La frontera se utiliza donde aporta testabilidad y reemplazo de tecnología; no se agregan microservicios ni un contenedor de inyección.')

page('Clases del núcleo y reglas de movimiento')
p('Los objetos de valor y las piezas describen el estado. Las estrategias describen cómo se mueve cada tipo de pieza. El validador reúne las condiciones necesarias para que una jugada completa sea legal.')
table(['Clase o interfaz','Responsabilidad'],[
('Position','Identifica una casilla mediante fila y columna; convierte coordenadas como e2 o e4. Es inmutable.'),
('Board','Guarda la matriz de 8 por 8, la casilla de captura al paso y el contador de medias jugadas. Permite copiar y restaurar el estado.'),
('Piece','Guarda tipo, color y marca de movimiento. Delega movimiento y ataque en una estrategia. Es inmutable.'),
('PieceColor','Representa blancas y negras, su dirección de avance y su fila inicial.'),
('PieceType','Identifica los seis tipos, sus personajes y sus valores. El rey tiene valor null, porque no posee puntaje numérico.'),
('PieceFactory','Selecciona y comparte las estrategias necesarias para crear piezas al comenzar y al promocionar.'),
('IMovementStrategy','Contrato de canMove y canAttack. Distingue avanzar de amenazar una casilla.'),
('PawnMovement y KnightMovement','El peón avanza y captura de forma distinta. El caballo comprueba el movimiento en L y puede saltar.'),
('RookMovement y BishopMovement','Comprueban movimientos rectos o diagonales y que no existan obstáculos.'),
('QueenMovement y KingMovement','La reina compone las estrategias de torre y alfil. El rey comprueba el movimiento de una casilla.'),
('CheckDetector','Busca el rey y consulta si una pieza rival amenaza su casilla.'),
('MoveValidator','Comprueba turno, límites, destino, movimiento, enroque y seguridad del rey mediante una simulación.')],[2.3,4.4])
h('Movimiento permitido y movimiento legal',2)
p('La estrategia de una torre puede permitir una jugada vertical. Sin embargo, esa jugada es ilegal si deja descubierto al propio rey. Por eso Piece comprueba la forma del movimiento y MoveValidator controla la legalidad completa. El enroque también requiere condiciones sobre el rey y la torre que exceden un movimiento geométrico.')
p('El valor del rey no se confunde con cero puntos: no puede ser capturado y perder por jaque mate termina la partida. Las capturas suman peón 1, caballo 3, alfil 3, torre 5 y reina 9.')

page('Clases de partida y sesión')
table(['Clase o interfaz','Responsabilidad'],[
('IGameService','Puerto de entrada del juego: mover, consultar estado y destinos, deshacer, rehacer y consultar puntos.'),
('Game','Coordina turnos, validación, historial, repetición, capturas, puntaje y resultado.'),
('Move y MoveResult','Move describe origen, destino y promoción. MoveResult indica si se aceptó la solicitud y comunica un mensaje.'),
('ICommand y MoveCommand','El contrato permite ejecutar y deshacer. El comando conserva una copia del tablero anterior para revertir los efectos.'),
('MoveRecord','Expone datos inmutables del historial. Los adaptadores no reciben comandos que puedan ejecutar libremente.'),
('IMatchService y MatchSession','Puerto y servicio de sesión: nueva partida, modo, intervención de IA, guardado y cronómetro.'),
('GameMode y Difficulty','Definen dos jugadores o solo, y cinco configuraciones de dificultad.'),
('SavedMatch','Datos del guardado: secuencia de movimientos, posición actual del historial, modo, nivel y marcas de tiempo.'),
('IGameRepository y PersistenceException','Contrato de carga y guardado; error que permite informar un fallo sin acoplar el núcleo a archivos.'),
('IChessAi y MinimaxAi','Contrato de elección de jugada y estrategia local que usa azar, capturas o búsqueda según el nivel.')],[2.35,4.35])
h('Historial y restauración',2)
p('MoveCommand conserva el estado necesario para restaurar capturas, promoción, enroque y al paso. Game mantiene las jugadas ejecutadas y las que pueden rehacerse. Una jugada nueva elimina la rama de rehacer. El puntaje se calcula desde las capturas del historial vigente, por lo que deshacer también resta puntos.')
p('Para guardar se registra la secuencia completa y el cursor del historial. Al cargar, MatchSession reproduce las jugadas con las mismas reglas y vuelve al cursor guardado. Esto recupera también derechos de enroque, al paso, repetición y posibilidad de rehacer.')
h('Tiempo y modo de juego',2)
p('MatchSession recibe un Clock por constructor. El tiempo cuenta desde la creación de la partida, incluso durante un cierre, y se detiene al terminar. En modo solo, el humano juega con blancas y la IA con negras. Deshacer normalmente revierte la jugada humana y la respuesta de la IA, para volver a la decisión anterior.')

page('Puertos de entrada y salida y adaptadores')
p('Un puerto es una interfaz que expresa qué operación está disponible o qué colaboración necesita el núcleo. Un adaptador resuelve la comunicación con una tecnología concreta. Entrada y salida describen quién inicia la interacción, no sólo si un método recibe o devuelve datos.')
table(['Puerto','Dirección','Función'],[
('IGameService','Entrada','La consola o la sesión solicitan jugadas y consultan el juego.'),
('IMatchService','Entrada','El adaptador web solicita operaciones sobre una sesión completa.'),
('IGameRepository','Salida','La sesión solicita cargar o guardar sin conocer archivos.'),
('IChessAi','Salida','La sesión pide una jugada sin depender del algoritmo elegido.')],[1.6,.9,4.2])
h('Adaptadores implementados',2)
p('WebServer utiliza el servidor HTTP incluido en el JDK. Traduce solicitudes como POST /api/move en objetos Move, llama a IMatchService y devuelve JSON. También sirve los archivos de la interfaz. Comprueba una revisión de estado para evitar aplicar solicitudes sobre una partida desactualizada.')
p('ConsoleUI recibe IGameService, lee texto como “e2 e4” y muestra el resultado en terminal. Demuestra que las reglas funcionan sin gráficos. Es una alternativa básica: no incorpora los controles web de tutorial, configuración de IA ni guardado de sesión.')
p('FileGameRepository implementa IGameRepository. Lee y escribe data/current-match.properties. Primero escribe un archivo temporal y luego reemplaza el guardado; utiliza un reemplazo atómico cuando el sistema de archivos lo permite. Un archivo inválido produce un error y no se sobrescribe silenciosamente.')
h('Ejemplo de entrada y salida en una captura',2)
for text in [
'1. El jugador selecciona e4 y luego d5. La interfaz ya mostró los destinos legales pedidos al servidor.',
'2. WebServer convierte la solicitud HTTP en un Move y llama a MatchSession.',
'3. La sesión verifica si corresponde jugar al humano y solicita a Game ejecutar la jugada.',
'4. MoveValidator consulta la estrategia, simula el movimiento y protege al rey. Game ejecuta el comando y cambia el turno.',
'5. MatchSession solicita guardar a IGameRepository. FileGameRepository escribe el archivo local.',
'6. El servidor devuelve el nuevo estado. La interfaz actualiza puntos, tablero e historial; si corresponde, solicita el turno de la IA.'
]:p(text)
p('Si falla el guardado, MatchSession restaura la sesión anterior y devuelve un rechazo. La interfaz no muestra una jugada como confirmada cuando no pudo conservarse.')

page('Patrones aplicados y verificación')
table(['Decisión','Aplicación concreta'],[
('Strategy','Piece recibe IMovementStrategy. La IA también se solicita a través de un contrato sustituible.'),
('Simple Factory','PieceFactory concentra la elección del comportamiento de cada tipo de pieza.'),
('Command','MoveCommand implementa execute y undo. Permite deshacer y rehacer efectos completos.'),
('Composición e inyección','La reina combina movimientos de torre y alfil. Los servicios reciben colaboradores por constructor.'),
('SOLID','Las responsabilidades se separan; las fronteras dependen de contratos y los datos internos se protegen con copias y vistas inmutables.')],[1.8,4.9])
h('Niveles de inteligencia artificial',2)
p('La IA es didáctica y local. No utiliza aprendizaje automático ni garantiza una fuerza profesional. Súper fácil elige al azar. Fácil prioriza capturas e introduce elecciones aleatorias. Los demás niveles comparan continuaciones mediante minimax con poda alfa-beta: dejan de explorar ramas que no pueden mejorar la decisión.')
table(['Nivel','Profundidad máxima','Límite de búsqueda'],[
('Medio','2 medias jugadas','1.200 nodos o 600 ms'),
('Difícil','3 medias jugadas','5.000 nodos o 1.200 ms'),
('Hiper mega difícil','4 medias jugadas','16.000 nodos o 2.200 ms')],[2.3,1.8,2.6])
p('Los límites son presupuestos internos, no un SLA. La búsqueda conserva la última profundidad terminada y usa el mismo validador que el jugador. En dos jugadores no se solicita ninguna jugada a la IA.')
h('Pruebas y reglas básicas',2)
p('Las 31 pruebas originales verifican estrategias, jaques, finales, jugadas especiales, restauración y 8.902 posiciones legales a profundidad tres. Otras 16 pruebas comprueban puntaje, ganador, cinco niveles, guardado, recuperación, reloj y fallo de escritura. Se ejecutan sin navegador con run.ps1 -Test.')
p('Los nuevos tests utilizan un repositorio en memoria y un reloj controlado: son dobles que permiten comprobar guardado y tiempo sin depender del disco o de esperar segundos reales. Una prueba adicional de esa suite utiliza el adaptador de archivo real.')
p('Puede haber mate con una sola pieza atacante si no existe defensa legal. Se verificó una torre con apoyo del rey y un caballo con el rey rival encerrado por sus propias piezas. No se transforma rey y caballo contra rey en un mate imposible.')
h('Relación con las clases',2)
p('Se aplican arquitectura de la unidad 00; SOLID, composición, patrones y pruebas de la unidad 01; y caché y búsqueda acotada de la unidad 04. Los contratos son interfaces de código, no un SLA de un proveedor externo.')
Path('build').mkdir(exist_ok=True)
for root in [doc.element,doc.styles.element]:
 for border in list(root.iter(qn('w:pBdr'))): border.getparent().remove(border)
doc.save('build/Informe-arquitectura-historico.docx')
print('Informe histórico creado en build/Informe-arquitectura-historico.docx; no reemplaza el informe vigente.')


